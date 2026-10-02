"""Real binary files, generated rather than committed.

The vendored corpus cannot cover these handlers because its data files are zero
bytes. These are genuine TIFF, HDF5, Office, PDF and zip payloads, built at test
time so nothing large or opaque lands in the repository.
"""

import io
import zipfile

__all__ = [
    "docx_bytes",
    "h5_bytes",
    "pdf_bytes",
    "pptx_bytes",
    "tiff_bytes",
    "xlsx_bytes",
    "zip_bytes",
]


def tiff_bytes(
    description: str = "SEM overview, 20 kV", size: tuple[int, int] = (8, 4)
) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    image = Image.new("L", size, color=128)
    image.save(buffer, format="TIFF", description=description)
    return buffer.getvalue()


def h5_bytes() -> bytes:
    import h5py

    buffer = io.BytesIO()
    with h5py.File(buffer, "w") as handle:
        handle.attrs["instrument"] = "SEM"
        group = handle.create_group("measurement")
        group.attrs["technique"] = "EDS"
        group.create_dataset("energy", data=[1.0, 2.0, 3.0])
        # Larger than inline_max_values, so only shape and dtype are recorded.
        group.create_dataset("counts", data=list(range(500)))
    return buffer.getvalue()


def docx_bytes(text: str = "Sample JM11 was measured at 20 kV.") -> bytes:
    import docx

    buffer = io.BytesIO()
    document = docx.Document()
    document.add_heading("Measurement report", level=1)
    document.add_paragraph(text)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "sample"
    table.cell(0, 1).text = "JM11"
    table.cell(1, 0).text = "technique"
    table.cell(1, 1).text = "EDS"
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    import openpyxl

    buffer = io.BytesIO()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "samples"
    sheet.append(["sample", "technique"])
    sheet.append(["JM11", "EDS"])
    book.save(buffer)
    return buffer.getvalue()


def pptx_bytes() -> bytes:
    import pptx

    buffer = io.BytesIO()
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Results"
    presentation.save(buffer)
    return buffer.getvalue()


def pdf_bytes(text: str = "Measurement report JM11") -> bytes:
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    page = canvas.Canvas(buffer)
    page.drawString(72, 720, text)
    page.save()
    return buffer.getvalue()


def zip_bytes(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return buffer.getvalue()
