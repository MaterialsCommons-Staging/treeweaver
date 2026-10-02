"""Office documents, in whichever of three shapes the caller asked for.

`ideas.txt` describes these as three handlers selected by an enabled flag. They
are one handler with a list of modes instead, because three handlers sharing one
pattern score the same specificity and which of them ran would come down to
registration order. A list also composes: `["keep", "text"]` puts the original
in the twin for an agent to open and the extracted text beside it for anything
that only reads text.
"""

import json
from typing import Literal

from pydantic import BaseModel, ConfigDict

from ..model import Copy, Handler, HandlerParam, Outcome, Write, meta_name

__all__ = ["OfficeSettings", "office_handler"]

Mode = Literal["keep", "hierarchy", "text"]


class OfficeSettings(BaseModel):
    model_config = ConfigDict(extra="forbid")

    modes: list[Mode] = ["keep"]


def _handle(param: HandlerParam) -> list[Outcome]:
    modes = list(getattr(param.settings, "modes", ["keep"]))
    path = param.node.path
    suffix = path.rsplit(".", 1)[-1].lower()

    outcomes: list[Outcome] = []
    structure = None

    for mode in modes:
        if mode == "keep":
            outcomes.append(Copy())
            continue
        if structure is None:
            with param.open() as stream:
                structure = _read(stream, suffix)
        if mode == "hierarchy":
            outcomes.append(
                Write(
                    dst=meta_name(path),
                    content=json.dumps(
                        structure, indent=2, sort_keys=True, default=str
                    ),
                )
            )
        elif mode == "text":
            outcomes.append(Write(dst=path + ".md", content=_as_text(structure)))
    return outcomes


def _read(stream, suffix: str) -> dict:
    if suffix == "docx":
        return _docx(stream)
    if suffix == "xlsx":
        return _xlsx(stream)
    if suffix == "pptx":
        return _pptx(stream)
    raise ValueError(f"no office reader for .{suffix}")


def _docx(stream) -> dict:
    import docx

    document = docx.Document(stream)
    return {
        "format": "docx",
        "properties": _core_properties(document.core_properties),
        "paragraphs": [
            {"style": p.style.name if p.style else None, "text": p.text}
            for p in document.paragraphs
            if p.text.strip()
        ],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ],
    }


def _xlsx(stream) -> dict:
    import openpyxl

    book = openpyxl.load_workbook(stream, data_only=True, read_only=True)
    sheets = {}
    for sheet in book.worksheets:
        rows = [
            [cell for cell in row]
            for row in sheet.iter_rows(values_only=True)
            if any(cell is not None for cell in row)
        ]
        sheets[sheet.title] = rows
    book.close()
    return {"format": "xlsx", "sheets": sheets}


def _pptx(stream) -> dict:
    import pptx

    presentation = pptx.Presentation(stream)
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        slides.append({"number": index, "text": texts})
    return {"format": "pptx", "slides": slides}


def _core_properties(properties) -> dict:
    keys = ("title", "author", "subject", "created", "modified", "last_modified_by")
    return {key: getattr(properties, key, None) for key in keys}


def _as_text(structure: dict) -> str:
    kind = structure.get("format")
    if kind == "docx":
        lines = [p["text"] for p in structure["paragraphs"]]
        for table in structure["tables"]:
            lines.extend(" | ".join(row) for row in table)
        return "\n\n".join(lines) + "\n"
    if kind == "xlsx":
        parts = []
        for name, rows in structure["sheets"].items():
            body = "\n".join(
                "\t".join("" if cell is None else str(cell) for cell in row)
                for row in rows
            )
            parts.append(f"## {name}\n\n{body}")
        return "\n\n".join(parts) + "\n"
    if kind == "pptx":
        parts = [
            f"## Slide {slide['number']}\n\n" + "\n\n".join(slide["text"])
            for slide in structure["slides"]
        ]
        return "\n\n".join(parts) + "\n"
    return ""


def office_handler() -> Handler:
    return Handler(
        name="office",
        match_regex=r"\.(docx|xlsx|pptx)$",
        handler_func=_handle,
        settings_model=OfficeSettings,
        settings=OfficeSettings(),
    )
